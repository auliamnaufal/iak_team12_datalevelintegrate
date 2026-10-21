import csv, json, glob
from openpyxl import load_workbook
import os
BASE=os.path.dirname(os.path.abspath(__file__))
U=os.path.join(BASE,'..','data_sumber')+os.sep
def rows_of(path):
    if path.endswith('.csv'):
        return list(csv.reader(open(path,encoding='utf-8-sig')))
    ws=load_workbook(path,read_only=True,data_only=True).active
    return [list(r) for r in ws.iter_rows(values_only=True)]
def num(x):
    try: return float(x)
    except: return None
def parse(path,keys):
    out={}
    for r in rows_of(path):
        if not r or not isinstance(r[0],str): continue
        v=[num(c) for c in r[1:1+len(keys)]]
        if r[0].strip().upper()==r[0].strip() and any(x is not None for x in v) and len(r[0])>3 and not r[0].startswith('38'):
            out[r[0].strip()]=dict(zip(keys,v))
    return out
def parse_api(path):
    j=json.load(open(path,encoding='utf-8'))
    vv=j['var'][0]['val']; th=j['tahun'][0]['val']; tt=j['turtahun'][0]['val']
    kmap={'Perkotaan':'pi_u','Perdesaan':'pi_r','Perkotaan+Perdesaan':'pi_t'}
    out={}
    for v in j['vervar']:
        rec={}
        for t in j['turvar']:
            k=kmap.get(t['label'])
            if k: rec[k]=j['datacontent'].get(f"{v['val']}{vv}{t['val']}{th}{tt}")
        if any(x is not None for x in rec.values()): out[v['label'].strip()]=rec
    return out
data={}
for n,v in parse_api(glob.glob(U+'bps_api_var656_*.json')[0]).items(): data.setdefault(n,{}).update(v)
src=[('Persentase_Rumah_Tangga_yang_Pernah*.csv',['rt_u','rt_r','rt_t']),
     ('Persentase_Penduduk_yang_Memiliki*.csv',['hp_u','hp_r','hp_t']),
     ('Persentase_Rumah_Tangga_yang_Memiliki_Menguasai_Komputer*.xlsx',['kom_u','kom_r','kom_t']),
     ('4_4_1*2024.csv',['tik'])]
for g,k in src:
    p=glob.glob(U+g)[0]; d=parse(p,k)
    for n,v in d.items(): data.setdefault(n,{}).update(v)
nat=data.pop('INDONESIA')
names=list(data)
rows=[dict(n=n.title().replace('Dki','DKI').replace('Di ','DI ').replace('Kep.','Kep.'),**{k:v for k,v in data[n].items()}) for n in names]
open(os.path.join(BASE,'..','data.js'),'w',encoding='utf-8').write('const D='+json.dumps({'rows':rows,'nat':nat},separators=(',',':'))+';\n')
print('data.js dibuat:',len(rows),'provinsi')
